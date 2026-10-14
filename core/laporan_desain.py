# core/laporan_desain.py
"""
Generator laporan siap cetak (DOCX dan PDF) dari satu daftar blok konten.

Blok:
  ("title", judul, subjudul)        ("h1", teks)        ("h2", teks)
  ("p", teks)   ("note", teks)      ("bullets", [teks])  ("kv", [(k, v)])
  ("table", {"header": [...], "rows": [[...]], "widths": [cm], "caption": s,
             "status_col": idx|None, "font": pt})
  ("image", path, lebar_cm, keterangan)      ("pagebreak",)

DOCX memakai python-docx; PDF memakai reportlab (font DejaVu dari paket
matplotlib agar karakter Unicode tampil). Tidak butuh LibreOffice.
"""

import os
from xml.sax.saxutils import escape

WARNA_STATUS = {
    "PASS": "C6EFCE", "WARNING": "FFEB9C", "ERROR": "FFC7CE",
    "CAPACITY EXCEEDANCE": "FF9999", "NOT ASSESSED": "D9D9D9",
}


# ------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------
def render_docx(blocks, path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for nm, pt in (("Heading 1", 15), ("Heading 2", 12)):
        doc.styles[nm].font.name = "Calibri"
        doc.styles[nm].font.size = Pt(pt)
        doc.styles[nm].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(90)
            r = p.add_run(b[1])
            r.bold = True
            r.font.size = Pt(24)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b[2])
            r.font.size = Pt(13)
        elif kind == "h1":
            doc.add_heading(b[1], level=1)
        elif kind == "h2":
            doc.add_heading(b[1], level=2)
        elif kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "note":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.italic = True
            r.font.size = Pt(8.5)
        elif kind == "bullets":
            for t in b[1]:
                doc.add_paragraph(t, style="List Bullet")
        elif kind == "kv":
            t = doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for k, v in b[1]:
                row = t.add_row().cells
                row[0].text, row[1].text = str(k), str(v)
                shade(row[0], "EAF0F7")
                row[0].paragraphs[0].runs[0].bold = True
                row[0].width, row[1].width = Cm(5.0), Cm(12.0)
            doc.add_paragraph()
        elif kind == "table":
            d = b[1]
            if d.get("caption"):
                p = doc.add_paragraph()
                r = p.add_run(d["caption"])
                r.bold = True
                r.font.size = Pt(9)
            t = doc.add_table(rows=1, cols=len(d["header"]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            fs = Pt(d.get("font", 8))
            for i, h in enumerate(d["header"]):
                c = t.rows[0].cells[i]
                c.text = str(h)
                shade(c, "1F3A5F")
                for r in c.paragraphs[0].runs:
                    r.bold = True
                    r.font.size = fs
                    r.font.color.rgb = RGBColor(255, 255, 255)
            for row in d["rows"]:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
                    for r in cells[i].paragraphs[0].runs:
                        r.font.size = fs
                    if d.get("status_col") == i and str(v) in WARNA_STATUS:
                        shade(cells[i], WARNA_STATUS[str(v)])
            if d.get("widths"):
                for row in t.rows:
                    for i, w in enumerate(d["widths"]):
                        row.cells[i].width = Cm(w)
            doc.add_paragraph()
        elif kind == "image":
            if os.path.exists(b[1]):
                doc.add_picture(b[1], width=Cm(b[2]))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(b[3])
                r.italic = True
                r.font.size = Pt(8.5)
        elif kind == "pagebreak":
            doc.add_page_break()
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)
    print("Laporan DOCX:", path)
    return path


# ------------------------------------------------------------------
# PDF
# ------------------------------------------------------------------
def render_pdf(blocks, path):
    import matplotlib
    from PIL import Image as PILImage
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (Image, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table,
                                    TableStyle, KeepTogether, ListFlowable,
                                    ListItem)

    fdir = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
    pdfmetrics.registerFont(TTFont("DV", os.path.join(fdir, "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DVB", os.path.join(
        fdir, "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DVI", os.path.join(
        fdir, "DejaVuSans-Oblique.ttf")))
    pdfmetrics.registerFontFamily("DV", normal="DV", bold="DVB", italic="DVI",
                                  boldItalic="DVB")
    ss = getSampleStyleSheet()
    normal = ParagraphStyle("n", parent=ss["Normal"], fontName="DV",
                            fontSize=9, leading=12.5, spaceAfter=4)
    note = ParagraphStyle("note", parent=normal, fontName="DVI", fontSize=7.5,
                          leading=10)
    h1 = ParagraphStyle("h1", parent=normal, fontName="DVB", fontSize=14,
                        leading=18, spaceBefore=10, spaceAfter=6,
                        textColor=colors.HexColor("#1F3A5F"))
    h2 = ParagraphStyle("h2", parent=h1, fontSize=11, leading=14,
                        spaceBefore=8)
    ttl = ParagraphStyle("t", parent=h1, fontSize=22, leading=28,
                         alignment=TA_CENTER, spaceBefore=150)
    sub = ParagraphStyle("s", parent=normal, fontSize=12, leading=16,
                         alignment=TA_CENTER)
    cap = ParagraphStyle("cap", parent=normal, fontName="DVB", fontSize=8,
                         spaceBefore=4, spaceAfter=2)
    figcap = ParagraphStyle("fc", parent=note, alignment=TA_CENTER)

    def cell(text, fs, bold=False, color=None):
        s = ParagraphStyle("c", parent=normal, fontName="DVB" if bold else "DV",
                           fontSize=fs, leading=fs + 2.2, spaceAfter=0,
                           textColor=color or colors.black)
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), s)

    lebar_hal = 21.0 * cm - 4.0 * cm
    story = []
    for b in blocks:
        kind = b[0]
        if kind == "title":
            story += [Paragraph(escape(b[1]), ttl), Spacer(1, 10),
                      Paragraph(escape(b[2]), sub), Spacer(1, 30)]
        elif kind == "h1":
            story.append(Paragraph(escape(b[1]), h1))
        elif kind == "h2":
            story.append(Paragraph(escape(b[1]), h2))
        elif kind == "p":
            story.append(Paragraph(escape(b[1]), normal))
        elif kind == "note":
            story.append(Paragraph(escape(b[1]), note))
        elif kind == "bullets":
            story.append(ListFlowable(
                [ListItem(Paragraph(escape(t), normal), leftIndent=12)
                 for t in b[1]], bulletType="bullet", start="•",
                leftIndent=12, bulletFontName="DV", bulletFontSize=8))
        elif kind == "kv":
            data = [[cell(k, 8.5, True), cell(v, 8.5)] for k, v in b[1]]
            t = Table(data, colWidths=[5.0 * cm, lebar_hal - 5.0 * cm])
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0F7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 8)]
        elif kind == "table":
            d = b[1]
            fs = d.get("font", 8)
            hdr = [cell(h, fs, True, colors.white) for h in d["header"]]
            data = [hdr] + [[cell(v, fs) for v in row] for row in d["rows"]]
            if d.get("widths"):
                tot = sum(d["widths"])
                cw = [w / tot * lebar_hal for w in d["widths"]]
            else:
                cw = [lebar_hal / len(d["header"])] * len(d["header"])
            t = Table(data, colWidths=cw, repeatRows=1)
            sty = [("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                   ("LEFTPADDING", (0, 0), (-1, -1), 3),
                   ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
            if d.get("status_col") is not None:
                for r, row in enumerate(d["rows"], start=1):
                    v = str(row[d["status_col"]])
                    if v in WARNA_STATUS:
                        sty.append(("BACKGROUND", (d["status_col"], r),
                                    (d["status_col"], r),
                                    colors.HexColor("#" + WARNA_STATUS[v])))
            t.setStyle(TableStyle(sty))
            if d.get("caption"):
                story.append(Paragraph(escape(d["caption"]), cap))
            story += [t, Spacer(1, 8)]
        elif kind == "image":
            if os.path.exists(b[1]):
                w, h = PILImage.open(b[1]).size
                lebar = min(b[2] * cm, lebar_hal)
                story.append(KeepTogether([
                    Image(b[1], width=lebar, height=lebar * h / w),
                    Paragraph(escape(b[3]), figcap), Spacer(1, 6)]))
        elif kind == "pagebreak":
            story.append(PageBreak())

    def nomor(canvas, doc):
        canvas.saveState()
        canvas.setFont("DV", 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawString(2 * cm, 1.1 * cm,
                          "Jembatan Baros - Laporan kewajaran dan tinjauan desain "
                          "(DRAFT, screening)")
        canvas.drawRightString(19 * cm, 1.1 * cm, "Hal. {}".format(doc.page))
        canvas.restoreState()

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm,
                      title="Laporan kewajaran dan tinjauan desain Jembatan "
                            "Baros").build(story, onFirstPage=nomor,
                                           onLaterPages=nomor)
    print("Laporan PDF:", path)
    return path
